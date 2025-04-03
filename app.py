import streamlit as st
import openai
import os
import faiss
from langchain_community.docstore.in_memory import InMemoryDocstore
from langchain_community.vectorstores import FAISS
from langchain_openai import OpenAIEmbeddings
from uuid import uuid4
from langchain_core.documents import Document
import json
from docx import Document
import re

# Configurazione API OpenAI
openai.api_key = st.secrets['api_key']
os.environ["OPENAI_API_KEY"] = openai.api_key
embeddings = OpenAIEmbeddings(model="text-embedding-3-large")

with open('contratti_template.json', 'r', encoding="utf-8") as file:
    contratti = json.load(file)

# Caricamento vector store
vector_store = FAISS.load_local('vector_store.faiss', embeddings, allow_dangerous_deserialization=True)


def save_contract_to_word(draft, filename="bozza_contratto.docx"):
    """Salva la bozza in un file Word all'interno della cartella 'contratti_generati'."""
    folder_path = "contratti_generati"
    os.makedirs(folder_path, exist_ok=True)  # Crea la cartella se non esiste
    file_path = os.path.join(folder_path, filename)

    doc = Document()
    doc.add_paragraph(draft)  # Aggiunge il testo della bozza
    doc.save(file_path)

    return file_path  # Ritorna il percorso del file salvato


# Funzione per generare una bozza basata sul template
def generate_contract_draft(template, user_desc):
    prompt = f"""Adatta il seguente contratto di una casa editrice alla richiesta dell'utente.

        **Template di riferimento:**
        {template["testo"]}

        **Descrizione richiesta dall'utente:**
        {user_desc}

        **Istruzioni importanti per l'adattamento:**
        - Se nella richiesta dell'utente sono presenti dettagli come titolo dell'opera, autore, percentuale di royalty o altre informazioni specifiche, incorporali nella bozza.
        - Se alcuni dati non sono forniti, lascia spazi segnaposto tra parentesi quadre, es: "[Titolo Opera]".
        - Mantieni il tono e la struttura del contratto originale.

        **Bozza generata:**
        """

    response = openai.chat.completions.create(
        model="gpt-4",
        messages=[{"role": "system", "content": "Sei un esperto di contratti editoriali."},
                  {"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def collect_missing_info(draft):
    # Trova le informazioni mancanti (quelle tra parentesi quadre)
    placeholders = re.findall(r'\[(.*?)\]', draft)
    return list(set(placeholders))  # Rimuove duplicati


def update_contract_with_params(draft, params):
    """Aggiorna il contratto sostituendo i parametri mancanti con i valori forniti dall'utente"""
    updated_draft = draft
    for key, value in params.items():
        if value:  # Se il valore è stato fornito
            updated_draft = re.sub(r'\[' + re.escape(key) + r'\]', value, updated_draft)
    return updated_draft


# Interfaccia Streamlit
def reset_chat():
    st.session_state.messages = []
    st.session_state.draft = None
    st.session_state.missing_info = []
    st.session_state.params = {}
    st.session_state.current_step = "input"  # Resetta al passaggio iniziale

    if "user_input" in st.session_state:
        del st.session_state["user_input"]
        st.session_state.user_input = None


def inizializza_stato():
    if "messages" not in st.session_state:
        st.session_state.messages = []

    if 'user_input' not in st.session_state:
        st.session_state.user_input = None

    if 'draft' not in st.session_state:
        st.session_state.draft = None

    if 'missing_info' not in st.session_state:
        st.session_state.missing_info = []

    if 'params' not in st.session_state:
        st.session_state.params = {}

    if 'current_step' not in st.session_state:
        st.session_state.current_step = "input"  # Può essere: "input", "review_draft", "fill_params", "final"

    if 'best_template' not in st.session_state:
        st.session_state.best_template = None


if "initialized" not in st.session_state:
    inizializza_stato()
    st.session_state.initialized = True  # Flag per evitare ripetizioni

with st.sidebar:
    st.image('giunti.jpg', width=400)
    st.markdown(
        """  
        ### Ciao! 👋  
        Il tuo assistente per la redazione di contratti d'autore ti dà il benvenuto! ✍️📖  

        🔹 **Come posso aiutarti oggi?**  
        Scrivimi qui a lato e descrivi il contratto che vuoi creare.  

        📌 **Cosa farò per te?**  
        - Troverò il **template più adatto** tra quelli disponibili.  
        - Ti proporrò **una bozza personalizzata**, pronta per essere perfezionata.  
        - Ti aiuterò a **completare tutti i parametri** mancanti.
        - Una volta perfezionata la bozza potrai **scaricarla in Word**.

        Sono qui per rendere il tuo lavoro più semplice e veloce! 😉  
        """
    )

    st.divider()

    if st.button("Riparti con una nuova richiesta"):
        reset_chat()

# Gestione dello stato dell'applicazione
if st.session_state.current_step == "input":
    # Fase iniziale: richiesta descrizione contratto
    if user_input := st.chat_input('Descrivi il contratto che desideri generare'):
        st.session_state.user_input = user_input
        st.session_state.messages.append({"role": "user", "content": user_input, "avatar": "user_icon.png"})

        with st.spinner("Sto elaborando la tua richiesta..."):
            # Trova il template più adatto
            best_template = vector_store.similarity_search(user_input, k=1)[0].metadata
            st.session_state.best_template = best_template

            # Genera la bozza
            draft = generate_contract_draft(best_template, user_input)
            st.session_state.draft = draft

            # Identifica i parametri mancanti
            missing_info = collect_missing_info(draft)
            st.session_state.missing_info = missing_info

            # Prepara il messaggio dell'assistente
            assistant_message = f"Ecco qua una proposta di bozza: \n\n{draft}\n\n"
            if missing_info:
                assistant_message += "\nSe vuoi **perfezionare il tuo contratto**, puoi darmi ulteriori dettagli su:\n\n"
                bullet_points = "\n".join([f"- *{key}*" for key in missing_info])
                assistant_message += bullet_points

            # Aggiorna lo stato
            st.session_state.messages.append(
                {"role": "assistant", "content": assistant_message, "avatar": 'assistant_icon.png'}
            )
            st.session_state.current_step = "fill_params"

        # Forza aggiornamento della pagina
        st.rerun()

elif st.session_state.current_step == "fill_params" and st.session_state.missing_info:
    # Mostra i messaggi precedenti
    for message in st.session_state.messages:
        with st.chat_message(message['role'], avatar=message.get('avatar', None)):
            st.write(message["content"])

    # Mostra informazioni sul template
    if st.session_state.best_template:
        st.markdown(
            f"**Template più simile trovato:** [{st.session_state.best_template['descrizione']}]({st.session_state.best_template['link']})")

    # Crea i campi di input per i parametri mancanti senza usare un form
    st.write("**Completa i parametri mancanti**")

    # Dizionario per memorizzare i valori dei parametri
    params = {}

    # Crea campi di input per ogni parametro mancante
    for param in st.session_state.missing_info:
        params[param] = st.text_input(f"{param}:", key=f"param_{param}")

    # Pulsante di aggiornamento
    if st.button("Aggiorna contratto"):
        # Salva i parametri inseriti
        st.session_state.params = params

        # Aggiorna il contratto con i parametri forniti
        updated_draft = update_contract_with_params(st.session_state.draft, params)
        st.session_state.draft = updated_draft

        # Verifica se ci sono ancora parametri mancanti
        remaining_missing = collect_missing_info(updated_draft)

        # Aggiorna il messaggio e lo stato
        user_message = "Ho compilato i parametri richiesti."
        st.session_state.messages.append({"role": "user", "content": user_message, "avatar": "user_icon.png"})

        if remaining_missing:
            # Ci sono ancora parametri da compilare
            assistant_message = f"Ho aggiornato la bozza con i parametri forniti. Ecco il risultato:\n\n{updated_draft}\n\n"
            assistant_message += "\nCi sono ancora alcuni parametri da compilare:\n\n"
            bullet_points = "\n".join([f"- *{key}*" for key in remaining_missing])
            assistant_message += bullet_points

            st.session_state.missing_info = remaining_missing
        else:
            # Tutti i parametri sono stati compilati
            assistant_message = f"Ecco il contratto completo:\n\n{updated_draft}\n\n"
            assistant_message += "\nIl contratto è stato completato con successo! Puoi scaricarlo come documento Word."

            # Genera il file Word
            file_path = save_contract_to_word(updated_draft)
            st.session_state.current_step = "final"

        st.session_state.messages.append(
            {"role": "assistant", "content": assistant_message, "avatar": 'assistant_icon.png'}
        )

        # Forza aggiornamento della pagina
        st.rerun()

elif st.session_state.current_step == "final":
    # Mostra i messaggi precedenti
    for message in st.session_state.messages:
        with st.chat_message(message['role'], avatar=message.get('avatar', None)):
            st.write(message["content"])

    # Mostra informazioni sul template
    if st.session_state.best_template:
        st.markdown(
            f"**Template più simile trovato:** [{st.session_state.best_template['descrizione']}]({st.session_state.best_template['link']})")

    # Genera il file Word quando si carica la pagina finale
    file_path = save_contract_to_word(st.session_state.draft)

    # Leggi il file per renderlo disponibile per il download
    with open(file_path, "rb") as file:
        file_content = file.read()

    # Mostra direttamente il pulsante di download
    st.download_button(
        label="Scarica contratto come Word",
        data=file_content,
        file_name="contratto_personalizzato.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

else:
    # Visualizza i messaggi nella fase corrente
    for message in st.session_state.messages:
        with st.chat_message(message['role'], avatar=message.get('avatar', None)):
            st.write(message["content"])

    # Mostra informazioni sul template se disponibile
    if st.session_state.best_template:
        st.markdown(
            f"**Template più simile trovato:** [{st.session_state.best_template['descrizione']}]({st.session_state.best_template['link']})")